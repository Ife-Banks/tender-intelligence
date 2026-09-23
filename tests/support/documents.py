"""Real document bytes for the prompt-09 suite, built in-process with the production libraries.

Fixtures are *generated* rather than checked in as binaries on purpose: a reviewer can read
exactly what each fixture contains, and the fixture cannot silently drift from the format the
extractor claims to handle. Each builder produces a genuine file of its format — a real PDF with
a real text layer, a real OOXML package, a real ZIP — so the extraction paths are exercised
end-to-end rather than against a mock that could disagree with the library.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pymupdf
from PIL import Image, ImageDraw

# Enough words for the language detector's MIN_TOKENS floor, in each supported language, and
# written the way a tender document actually reads (function words plus domain nouns).
ENGLISH = (
    "The World Health Organization invites interested consulting firms to submit proposals for "
    "the development of a regional regulatory framework. The successful bidder shall provide a "
    "detailed methodology, a work plan and evidence of experience in the health sector. All "
    "documents must be submitted in English and must be received before the deadline. The "
    "contract will be awarded to the firm whose proposal best meets the evaluation criteria."
)

FRENCH = (
    "L'Organisation mondiale de la sante invite les cabinets de conseil interesses a soumettre "
    "des propositions pour le developpement d'un cadre reglementaire regional. Le soumissionnaire "
    "retenu doit fournir une methodologie detaillee, un plan de travail et la preuve de son "
    "experience dans le secteur de la sante. Tous les documents doivent etre soumis en francais "
    "et doivent etre recus avant la date limite. Le contrat sera attribue au cabinet dont la "
    "proposition repond le mieux aux criteres d'evaluation."
)

PORTUGUESE = (
    "A Organizacao Mundial da Saude convida as empresas de consultoria interessadas a apresentar "
    "propostas para o desenvolvimento de um quadro regulamentar regional. O concorrente "
    "selecionado deve fornecer uma metodologia detalhada, um plano de trabalho e prova da sua "
    "experiencia no setor da saude. Todos os documentos devem ser apresentados em portugues e "
    "devem ser recebidos antes do prazo. O contrato sera atribuido a empresa cuja proposta "
    "responde melhor aos criterios de avaliacao."
)

#: An evaluation matrix shaped like the ones tenders actually publish — the table prompt 09 §9
#: exists to preserve, because a later verdict stage reads weights and thresholds from it.
EVALUATION_MATRIX = [
    ["Criterion", "Weight", "Threshold"],
    ["Technical methodology", "40%", "70"],
    ["Team qualifications", "30%", "65"],
    ["Financial proposal", "20%", "60"],
    ["Regional experience", "10%", "50"],
]

#: A deadline table: the single most business-critical fact in most tender documents.
DEADLINE_TABLE = [
    ["Milestone", "Date", "Responsible"],
    ["Clarification questions due", "12 March 2026", "Bidder"],
    ["Proposal submission deadline", "30 March 2026", "Bidder"],
    ["Contract award notification", "24 April 2026", "Contracting authority"],
]

#: An eligibility table — the third shape prompt 09 §20 names (deadline, eligibility, evaluation
#: matrix). Rows are requirements, so a mangled row means a bidder misreads what disqualifies them.
ELIGIBILITY_TABLE = [
    ["Requirement", "Mandatory", "Evidence required"],
    ["Registered in a WHO member state", "Yes", "Certificate of incorporation"],
    ["Minimum five years in health consulting", "Yes", "Three reference contracts"],
    ["Annual turnover above USD 500,000", "No", "Audited financial statements"],
]


def text_pdf(pages: list[str]) -> bytes:
    """A PDF with a genuine native text layer, one entry of *pages* per page."""
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(56, 56, 540, 780), text, fontsize=11, fontname="helv")
    data: bytes = document.tobytes()
    document.close()
    return data


def blank_pdf(page_count: int = 1) -> bytes:
    """A page with no text at all: the degenerate case of "needs OCR"."""
    document = pymupdf.open()
    for _ in range(page_count):
        document.new_page()
    data: bytes = document.tobytes()
    document.close()
    return data


def image_pdf(text: str = "SCANNED ANNEX", page_count: int = 1) -> bytes:
    """An image-only PDF: what a scanned annex looks like to a text extractor.

    This is exactly the docs/06 §6.2 "scanned PDF (image) → OCR" input. ``text`` is drawn onto the
    raster so the fixture is a believable scan, but the page carries no text layer, so extraction
    can only succeed through OCR.
    """
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 200), text, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    document = pymupdf.open()
    for _ in range(page_count):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=buffer.getvalue())
    data: bytes = document.tobytes()
    document.close()
    return data


def mixed_pdf(native_pages: int = 1, scanned_pages: int = 1) -> bytes:
    """A native-text RFP with a scanned annex appended — the realistic mixed case (§4.1/§4.2).

    The per-page decision is what makes this fixture matter: a document-level choice would either
    OCR the native pages or drop the scanned ones, and both are wrong.
    """
    document = pymupdf.open()
    for _ in range(native_pages):
        page = document.new_page()
        page.insert_textbox(
            pymupdf.Rect(56, 56, 540, 780), ENGLISH, fontsize=11, fontname="helv"
        )

    image = Image.new("RGB", (1240, 1754), "white")
    ImageDraw.Draw(image).text((100, 200), "SCANNED ANNEX", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    for _ in range(scanned_pages):
        page = document.new_page()
        page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=buffer.getvalue())

    data: bytes = document.tobytes()
    document.close()
    return data


def table_pdf(
    rows: list[list[str]],
    *,
    caption: str = "",
    top: float = 120.0,
    left: float = 56.0,
    right_margin: float = 56.0,
    page_width: float = 595.0,
    font_size: float = 9.0,
    padding: float = 6.0,
    min_column_width: float = 60.0,
) -> bytes:
    """A PDF containing a ruled-line table, drawn the way real tender annexes draw them.

    Ruled lines (not whitespace alignment) are used deliberately: ``strategy="lines"`` detection
    keys off them, and a whitespace-aligned fixture would test a path the production code does not
    take.

    Column widths are measured from the widest cell in each column rather than fixed, because
    ``insert_textbox`` draws nothing at all when its text does not fit its box. A fixed width would
    therefore silently omit exactly the long, information-bearing cells — a requirement, an
    evidence clause — that a table fixture exists to check, and the resulting test would pass while
    asserting against a table that never contained the text.
    """
    document = pymupdf.open()
    page = document.new_page()
    if caption:
        page.insert_textbox(
            pymupdf.Rect(left, 56, 540, top), caption, fontsize=11, fontname="helv"
        )

    columns = max(len(row) for row in rows)
    widths = [
        max(
            min_column_width,
            max(
                (
                    pymupdf.get_text_length(row[index], fontname="helv", fontsize=font_size)
                    for row in rows
                    if index < len(row)
                ),
                default=0.0,
            )
            + 2 * padding,
        )
        for index in range(columns)
    ]
    available = page_width - left - right_margin
    if sum(widths) > available:
        widths = [width * available / sum(widths) for width in widths]

    row_height = font_size * 1.6 + 2 * padding
    height = row_height * len(rows)
    bottom = top + height

    edges = [left]
    for width in widths:
        edges.append(edges[-1] + width)

    for index in range(len(rows) + 1):
        y = top + row_height * index
        page.draw_line(pymupdf.Point(left, y), pymupdf.Point(edges[-1], y))
    for x in edges:
        page.draw_line(pymupdf.Point(x, top), pymupdf.Point(x, bottom))

    for row_index, row in enumerate(rows):
        for cell_index, cell in enumerate(row):
            page.insert_textbox(
                pymupdf.Rect(
                    edges[cell_index] + padding / 2,
                    top + row_height * row_index + padding / 2,
                    edges[cell_index + 1] - padding / 2,
                    top + row_height * (row_index + 1) - padding / 2,
                ),
                cell,
                fontsize=font_size,
                fontname="helv",
            )

    data: bytes = document.tobytes()
    document.close()
    return data


def docx_bytes(blocks: list[str | list[list[str]]]) -> bytes:
    """A DOCX whose *blocks* are paragraphs (``str``) or tables (list of rows), in that order."""
    document = docx.Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
            continue
        table = document.add_table(rows=len(block), cols=len(block[0]))
        for row_index, row in enumerate(block):
            for cell_index, cell in enumerate(row):
                table.cell(row_index, cell_index).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def zip_bytes(members: dict[str, bytes]) -> bytes:
    """A ZIP container holding *members*, keyed by member name."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return buffer.getvalue()
