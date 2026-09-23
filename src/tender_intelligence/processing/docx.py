""":mod:`tender_intelligence.processing.docx` — DOCX extraction (prompt 09 §6).

WAHO publishes DOCX TORs (docs/06 §6.2), so DOCX is a first-class input. Extraction uses
python-docx, which reads the OOXML package directly — Word/Office is never required.
The document is read from bytes, so no temporary files are involved.

Structure: ``iter_inner_content()`` yields paragraphs and tables in true document order. The
separate ``.paragraphs`` and ``.tables`` collections do not preserve that interleaving, so a
deadline paragraph that follows an evaluation matrix would be indistinguishable from one that
precedes it. Each block becomes an :class:`ExtractedSection`, which is the "appropriate structural
representation" prompt 09 §8 requires in place of inventing fake PDF page numbers for a format
that has no pages.

Merged cells: python-docx repeats a merged cell's text for each grid position it spans. That is
recorded as-is rather than de-duplicated, because collapsing spans would misreport the table's
true column structure.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

import docx
from docx.opc.exceptions import OpcError, PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from tender_intelligence.processing.errors import ParseFailedError
from tender_intelligence.processing.representation import ExtractedSection, ExtractedTable

_METHOD = "docx"


@dataclass(frozen=True)
class DocxExtractionResult:
    """Extracted DOCX content as ordered sections plus a flat table/reading-order view."""

    sections: list[ExtractedSection] = field(default_factory=list)
    tables: list[ExtractedTable] = field(default_factory=list)
    text: str = ""
    method: str = _METHOD


def _table_from_docx(table: Table, *, section_index: int, table_index: int) -> ExtractedTable:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(row)]
    has_header = len(rows) > 1
    return ExtractedTable(
        location=f"section:{section_index}",
        index=table_index,
        headers=list(rows[0]) if has_header else [],
        rows=rows,
        header_convention="first_row" if has_header else "none",
        page=None,
    )


def extract_docx(data: bytes) -> DocxExtractionResult:
    """Extract *data* as a DOCX.

    Raises :class:`ParseFailedError` when the bytes are not a readable DOCX package.
    """
    try:
        document = docx.Document(io.BytesIO(data))
    except (
        zipfile.BadZipFile,
        PackageNotFoundError,
        OpcError,
        ValueError,
        KeyError,
    ) as exc:
        raise ParseFailedError(
            f"docx could not be parsed: {type(exc).__name__}",
            context={"format": "docx", "exception": type(exc).__name__},
        ) from exc

    sections: list[ExtractedSection] = []
    tables: list[ExtractedTable] = []
    parts: list[str] = []

    try:
        blocks = list(document.iter_inner_content())
    except Exception as exc:
        raise ParseFailedError(
            f"docx body could not be read: {type(exc).__name__}",
            context={"format": "docx", "exception": type(exc).__name__},
        ) from exc

    for index, block in enumerate(blocks):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            sections.append(ExtractedSection(index=index, kind="paragraph", text=text))
            parts.append(text)
            continue

        if isinstance(block, Table):
            table = _table_from_docx(block, section_index=index, table_index=len(tables))
            if not table.rows:
                continue
            tables.append(table)
            # The table's textual rendering is carried on its section so reading order is intact;
            # the structured rows remain available on both the section and the document.
            sections.append(
                ExtractedSection(
                    index=index, kind="table", text=table.to_text(), tables=[table]
                )
            )
            parts.append(table.to_text())

    return DocxExtractionResult(
        sections=sections,
        tables=tables,
        text="\n\n".join(parts).strip(),
        method=_METHOD,
    )
